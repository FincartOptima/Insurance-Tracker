"""Excel reports built from exactly the same filtered result as the dashboard."""
import datetime as dt
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("client_name", "Client name", 28), ("client_email", "Email", 34),
    ("phone", "Phone", 20), ("rm_name", "Relationship manager", 25), ("team", "Team", 20),
    ("policy_no", "Policy number", 26), ("policy", "Policy", 40),
    ("policy_partner", "Insurer", 28), ("insurance_type", "Insurance type", 22),
    ("sum_assured", "Sum insured (INR)", 22), ("premium_amount", "Premium (INR)", 22),
    ("next_premium_date", "Renewal date", 19), ("days_until", "Days to renewal", 19),
    ("status", "Follow-up status", 20), ("remarks", "Remarks", 50),
    ("rescheduled", "Rescheduled", 16),
]


def _set_value(cell, value):
    # Text stays literal, including leading zeros, +phone numbers and =notes.
    # Never let data imported from a client workbook become an Excel formula.
    if isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
        cell.data_type = "s"
    else:
        cell.value = value


def workbook(data):
    wb = Workbook()
    ws = wb.active
    ws.title = "Policies"
    ws.append([label for _, label, _ in COLUMNS])
    for row_index, row in enumerate(data["rows"], 2):
        for col_index, (key, _, _) in enumerate(COLUMNS, 1):
            value = row.get(key)
            if key == "next_premium_date" and value:
                value = dt.date.fromisoformat(value)
            if key == "rescheduled":
                value = "Yes" if value else "No"
            cell = ws.cell(row_index, col_index)
            _set_value(cell, value)
            if key in ("premium_amount", "sum_assured"):
                cell.number_format = '#,##0.00'
            elif key == "next_premium_date":
                cell.number_format = "dd mmm yyyy"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for index, (_, _, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    meta = wb.create_sheet("Applied filters")
    meta.append(["Filter", "Value"])
    for label, value in [
        ("As of (Asia/Kolkata)", data["as_of"]), ("Timeline", data["bucket_label"]),
        ("Start date (inclusive)", data["start"] or ""), ("End date (inclusive)", data["end"] or ""),
        ("Team", data["team"]), ("Relationship manager", data["rm"]),
        ("Insurance type", data["insurance_type"]), ("Follow-up status", data["status"]),
        ("Search", data["search"]), ("Matching policies", len(data["rows"])),
        ("Total premium (INR)", data["premium_total"]),
        ("Timeline rule", "Next N days includes today through N days ahead. All upcoming excludes overdue and undated policies."),
    ]:
        meta.append([label, None])
        _set_value(meta.cell(meta.max_row, 2), value)
    meta.column_dimensions["A"].width = 29
    meta.column_dimensions["B"].width = 95
    meta.freeze_panes = "A2"
    for sheet in wb:
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="234E70")
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.alignment = Alignment(vertical="center")
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
