"""Parse an Insurance export and upsert renewal records into Postgres."""
import datetime as dt
import io
import os
import re

from openpyxl import load_workbook

import db

COLUMN_MAP = {
    "PolicyNo": "policy_no",
    "ClientName": "client_name",
    "ClientEmail": "client_email",
    "RmName": "rm_name",
    "Policy": "policy",
    "PolicyPartner": "policy_partner",
    "InsuranceType": "insurance_type",
    "SumAssured": "sum_assured",
    "PremiumAmount": "premium_amount",
    "NextPremimum_Date": "next_premium_date",
    "recordStatus": "record_status",
}
NUMERIC_COLS = {"sum_assured", "premium_amount"}

UNASSIGNED_TEAM = "Unassigned"


class IngestError(Exception):
    pass


def _norm_name(value):
    if value in (None, ""):
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().casefold()


def _conflict_update(column):
    """Upsert clause for one column on a re-uploaded policy.

    next_premium_date is special: once someone has manually rescheduled a
    policy (rescheduled_at is set), a later re-upload of the same source row
    must not silently overwrite that with the file's original, stale date.
    """
    if column == "next_premium_date":
        return (
            "next_premium_date = CASE WHEN renewals.rescheduled_at IS NULL "
            "THEN EXCLUDED.next_premium_date ELSE renewals.next_premium_date END"
        )
    return f"{column} = EXCLUDED.{column}"


def _open_workbook(path):
    """Open .xlsx, or a .xls that is really a renamed .xlsx (these exports are).

    openpyxl refuses a path ending in .xls before it ever looks at the bytes,
    so the content is handed over as a stream instead.
    """
    with open(path, "rb") as fh:
        payload = fh.read()
    if payload[:2] != b"PK":
        raise IngestError(
            "This file is a genuine legacy .xls, which this app cannot read. "
            "Please open it in Excel, re-save as .xlsx, and upload again."
        )
    return load_workbook(io.BytesIO(payload), data_only=True, read_only=True)


def _to_float(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return None if value != value else float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def _to_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d-%m-%Y %H:%M:%S",
                "%d-%m-%Y", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def normalize_whatsapp_number(value):
    """Digits-only, country-code-prefixed number for a wa.me link, or None."""
    if value in (None, ""):
        return None
    digits = re.sub(r"\D", "", str(value))
    if not digits:
        return None
    if len(digits) == 10:
        return "91" + digits  # bare Indian mobile number
    if len(digits) == 11 and digits.startswith("0"):
        return "91" + digits[1:]  # leading trunk 0
    return digits  # already carries a country code (or unusual) - use as-is


def resolve_team(rm_name, emp_map):
    """RM name -> Team. Exact match first, then a subset-of-tokens match
    (so 'Swaraj Thakur' resolves against an employee record of
    'Swaraj Singh Thakur'). Blank or genuinely unmatched -> Unassigned,
    with the raw name returned for reporting in that second case.
    """
    name = _norm_name(rm_name)
    if not name:
        return UNASSIGNED_TEAM, None
    if name in emp_map:
        return emp_map[name], None

    tokens = set(name.split())
    hits = [full for full in emp_map if tokens <= set(full.split())]
    if len(hits) == 1:
        return emp_map[hits[0]], None
    return UNASSIGNED_TEAM, str(rm_name).strip()


def recompute_teams(conn):
    """Re-resolve team for every stored policy from the current employee_ref.

    Always applies to every row (not just ones missing a team) - rm_name
    itself gets refreshed on every insurance re-upload, so a policy that
    changed hands to a different RM should have its team follow.
    """
    with conn.cursor() as cur:
        cur.execute("SELECT lower(name) AS name, team FROM employee_ref")
        emp_map = {r["name"]: r["team"] for r in cur.fetchall()}

        cur.execute("SELECT policy_no, rm_name FROM renewals")
        rows = cur.fetchall()

        unmatched = set()
        for r in rows:
            team, unmatched_name = resolve_team(r["rm_name"], emp_map)
            if unmatched_name:
                unmatched.add(unmatched_name)
            cur.execute("UPDATE renewals SET team = %s WHERE policy_no = %s",
                        (team, r["policy_no"]))
    conn.commit()
    return sorted(unmatched)


def ingest_employee_file(path, source_name):
    """Load an Employee Reference file (Emp Code, Team, Name) and replace
    the stored mapping, then re-resolve every policy's team against it.
    """
    db.init_db()
    conn = db.connect()
    try:
        wb = _open_workbook(path)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        try:
            header = [str(c).strip() if c is not None else "" for c in next(rows)]
        except StopIteration:
            raise IngestError("The employee reference file is empty.")

        norm = [h.strip().casefold() for h in header]

        def find(*keywords):
            for i, h in enumerate(norm):
                if any(k in h for k in keywords):
                    return i
            return None

        code_i = find("emp code", "empcode", "employee code")
        team_i = find("team")
        name_i = find("name")
        if team_i is None or name_i is None:
            raise IngestError("Couldn't find both a Team and a Name column in this file.")

        entries = []
        for raw in rows:
            if raw is None or all(v is None for v in raw):
                continue
            team = raw[team_i] if team_i < len(raw) else None
            name = raw[name_i] if name_i < len(raw) else None
            if not team or not name:
                continue
            code = raw[code_i] if (code_i is not None and code_i < len(raw)) else None
            entries.append((
                str(code).strip() if code else None,
                str(team).strip(),
                str(name).strip(),
            ))
        wb.close()

        if not entries:
            raise IngestError("No usable rows (with both a Team and a Name) found.")

        with conn.cursor() as cur:
            cur.execute("DELETE FROM employee_ref")
            cur.executemany(
                "INSERT INTO employee_ref (emp_code, team, name) VALUES (%s, %s, %s)",
                entries,
            )
        conn.commit()

        unmatched = recompute_teams(conn)
        db.set_meta(conn, "last_employee_upload", dt.datetime.now().isoformat(timespec="seconds"))
        db.set_meta(conn, "last_employee_source", source_name)
        conn.commit()

        return {"loaded": len(entries), "unmatched": unmatched}
    finally:
        conn.close()


def ingest_file(insurance_path, source_name):
    """Load one export into the DB. Returns a summary dict for the UI."""
    db.init_db()
    conn = db.connect()
    try:
        wb = _open_workbook(insurance_path)
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        try:
            header = [str(c).strip() if c is not None else "" for c in next(rows)]
        except StopIteration:
            raise IngestError("The uploaded file is empty.")

        idx = {h: i for i, h in enumerate(header)}
        missing = [h for h in ("recordStatus", "PolicyNo", "NextPremimum_Date")
                   if h not in idx]
        if missing:
            raise IngestError(
                "This does not look like an Insurance export - missing column(s): "
                + ", ".join(missing)
            )

        now = dt.datetime.now().isoformat(timespec="seconds")
        inserted = updated = skipped = no_date = 0
        seen_premium, conflicts = {}, []

        with conn.cursor() as cur:
            for raw in rows:
                if raw is None or all(v is None for v in raw):
                    continue
                rec = {}
                for src, dest in COLUMN_MAP.items():
                    i = idx.get(src)
                    rec[dest] = raw[i] if (i is not None and i < len(raw)) else None

                if str(rec.get("record_status") or "").strip().casefold() != "confirmed":
                    skipped += 1
                    continue

                policy_no = str(rec.get("policy_no") or "").strip()
                if not policy_no:
                    skipped += 1
                    continue

                for col in NUMERIC_COLS:
                    rec[col] = _to_float(rec[col])
                # Trimmed so a stray trailing space doesn't quietly split one
                # RM into two entries in the RM filter dropdown.
                if rec["rm_name"]:
                    rec["rm_name"] = str(rec["rm_name"]).strip()
                rec["next_premium_date"] = _to_date(rec["next_premium_date"])
                if rec["next_premium_date"] is None:
                    no_date += 1

                # Same PolicyNo twice in one file with a different premium is a
                # data conflict, not a re-upload: the last row wins either way,
                # so it has to be surfaced rather than silently overwritten.
                if policy_no in seen_premium:
                    prior = seen_premium[policy_no]
                    if prior != rec["premium_amount"]:
                        conflicts.append({
                            "policy_no": policy_no,
                            "client": rec["client_name"],
                            "policy": rec["policy"],
                            "kept": rec["premium_amount"],
                            "dropped": prior,
                        })
                seen_premium[policy_no] = rec["premium_amount"]

                rec.pop("record_status", None)
                rec["policy_no"] = policy_no
                rec["source_file"] = source_name
                rec["uploaded_at"] = now

                cur.execute("SELECT 1 FROM renewals WHERE policy_no=%s", (policy_no,))
                exists = cur.fetchone() is not None

                cols = list(rec)
                cur.execute(
                    "INSERT INTO renewals ({}) VALUES ({}) "
                    "ON CONFLICT (policy_no) DO UPDATE SET {}".format(
                        ", ".join(cols),
                        ", ".join(["%s"] * len(cols)),
                        ", ".join(_conflict_update(c) for c in cols if c != "policy_no"),
                    ),
                    [rec[c] for c in cols],
                )
                if exists:
                    updated += 1
                else:
                    inserted += 1

        wb.close()
        db.set_meta(conn, "last_upload", now)
        db.set_meta(conn, "last_source", source_name)
        conn.commit()
        unmatched_rms = recompute_teams(conn)

        return {
            "inserted": inserted,
            "updated": updated,
            "skipped": skipped,
            "no_date": no_date,
            "total_rows": db.count_renewals(conn),
            "conflicts": conflicts,
            "unmatched_rms": unmatched_rms,
        }
    finally:
        conn.close()
