import datetime as dt
import unittest
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

import app
import exports
import queries

TODAY = dt.date(2026, 9, 27)


def policy(number, days, **overrides):
    r = dict(policy_no=number, client_name='Sample Client', client_email='client@example.com',
             phone='09876543210', rm_name='Riya', team='North', policy='Health plan',
             policy_partner='Example Insurance', insurance_type='Health',
             sum_assured=Decimal('500000.00'), premium_amount=Decimal('12345.67'),
             next_premium_date=TODAY + dt.timedelta(days=days) if days is not None else None,
             status='Not Done', remarks='', rescheduled_at=None)
    return {**r, **overrides}


class Connection:
    def __init__(self, rows):
        self.rows = rows
        self.calls = []
        self.rowcount = 1
        self.commits = self.rollbacks = self.closed = 0
    def cursor(self): return self
    def __enter__(self): return self
    def __exit__(self, *args): pass
    def execute(self, sql, params=None): self.calls.append((sql, params))
    def fetchall(self): return self.rows
    def fetchone(self): return {'c': len(self.rows)}
    def close(self): self.closed += 1
    def commit(self): self.commits += 1
    def rollback(self): self.rollbacks += 1


class RenewalTests(unittest.TestCase):
    def setUp(self):
        self.rows = [policy(str(d), d) for d in [-1, 0, 1, 2, 7, 8, 14, 30, 60, 90, 91, None]]
        self.conn = Connection(self.rows)
    def data(self, **kwargs): return queries.build(self.conn, today=TODAY, **kwargs)
    def test_timeline_boundaries(self):
        expected = {'overdue':1,'today':1,'tomorrow':2,'next2':3,'next7':4,
                    'next14':6,'next30':7,'next60':8,'next90':9,'upcoming':10,'all':12,'no_date':1}
        for bucket, count in expected.items():
            with self.subTest(bucket=bucket):
                self.assertEqual(len(self.data(bucket=bucket)['rows']), count)
    def test_custom_includes_both_endpoints(self):
        data = self.data(bucket='custom', start='2026-09-28', end='2026-10-04')
        self.assertEqual([r['days_until'] for r in data['rows']], [1,2,7])
        self.assertEqual(len(self.data(bucket='custom',start='2026-09-27',end='2026-09-27')['rows']),1)
    def test_invalid_filters(self):
        for kw in [dict(bucket='oops'),dict(status='bad'),dict(bucket='custom'),
                   dict(bucket='custom',start='bad',end='2026-10-01'),
                   dict(bucket='custom',start='2026-10-10',end='2026-09-01')]:
            with self.assertRaises(ValueError): self.data(**kw)
    def test_combined_filters_and_counts(self):
        self.conn.rows = [policy('00042', 7, client_name='Asha', remarks='=1+1'),
                          policy('00043', 7, team='South'),policy('00044',7,status='Done'),
                          policy('00045',7,insurance_type='Life'),policy('00046',7,rm_name='Neha')]
        data = self.data(bucket='next7', search=' asha ', insurance_type='Health',team='North',rm='Riya',status='Not Done')
        self.assertEqual([r['policy_no'] for r in data['rows']],['00042'])
        self.assertEqual(data['counts']['next7'],1)
        self.assertEqual(data['team_options'],['North','South'])
    def test_search_contact_and_policy(self):
        self.conn.rows=[policy('00042',1)]
        for search in ['00042','CLIENT@','987654','sample']:
            self.assertEqual(len(self.data(search=search)['rows']),1)
    def test_unassigned_and_missing_date(self):
        self.conn.rows=[policy('0',None,team=None)]
        self.assertEqual(len(self.data(bucket='all',team='Unassigned')['rows']),1)
        self.assertEqual(self.data()['no_date'],1)
    def test_excel_types_complete_fields_and_literal_text(self):
        self.conn.rows=[policy('00042',1,client_name='=DANGEROUS()',remarks='=1+1')]
        wb=load_workbook(exports.workbook(self.data()))
        ws=wb['Policies']
        self.assertEqual(ws.max_column,16)
        self.assertEqual(ws['A2'].data_type,'s')
        self.assertEqual(ws['O2'].data_type,'s')
        self.assertEqual(ws['F2'].value,'00042')
        self.assertEqual(ws['C2'].value,'09876543210')
        self.assertEqual(ws['K2'].value,12345.67)
        self.assertEqual(ws['L2'].value.date(),TODAY+dt.timedelta(days=1))
        self.assertEqual(ws.freeze_panes,'C2')
        self.assertEqual(wb.sheetnames,['Policies','Applied filters'])
    def test_empty_export_has_headers(self):
        self.conn.rows=[]
        ws=load_workbook(exports.workbook(self.data())).active
        self.assertEqual(ws.max_row,1)
        self.assertEqual(ws.max_column,16)
    def test_api_export_parity_and_more_than_one_page(self):
        self.conn.rows=[policy(f'000{i:03}',7) for i in range(60)]+[policy('excluded',30)]
        with patch.object(app.db,'connect',return_value=self.conn), patch.object(queries,'business_today',return_value=TODAY):
            client=app.app.test_client()
            args='?bucket=next7&team=North&rm=Riya&type=Health&status=Not+Done&q=Sample'
            response=client.get('/api/renewals'+args)
            export=client.get('/api/renewals/export'+args)
            self.assertEqual(response.status_code,200)
            self.assertEqual(export.status_code,200)
            ws=load_workbook(BytesIO(export.data)).active
            self.assertEqual(ws.max_row-1,60)
            self.assertEqual([r[5] for r in list(ws.values)[1:]], [r['policy_no'] for r in response.json['rows']])
            self.assertEqual(export.headers['Cache-Control'],'no-store')
            for route in ['/api/renewals','/api/renewals/export']:
                self.assertEqual(client.get(route+'?bucket=custom&start=2026-10-01&end=2026-09-01').status_code,400)
    def test_atomic_validation_before_writes(self):
        with self.assertRaises(ValueError): queries.update_fields(self.conn,'00042',{'phone':'123','status':'bad'})
        self.assertEqual(self.conn.calls,[])
        self.assertEqual(self.conn.commits,0)
    def test_atomic_save_and_slash_policy(self):
        with patch.object(app.db,'connect',return_value=self.conn):
            response=app.app.test_client().patch('/api/renewals/HI%2F2026%2F001',json={'fields':{'phone':'0123','remarks':'Contacted','status':'Done'}})
        self.assertEqual(response.status_code,200)
        self.assertEqual(self.conn.commits,1)
        self.assertTrue(all(params[-1]=='HI/2026/001' for _,params in self.conn.calls))
    def test_missing_policy_rollback(self):
        self.conn.rowcount=0
        with self.assertRaises(ValueError): queries.update_fields(self.conn,'missing',{'status':'Done'})
        self.assertEqual(self.conn.rollbacks,1)
        self.assertEqual(self.conn.commits,0)
    def test_templates_render(self):
        with app.app.test_request_context('/'):
            html=app.render_template('dashboard.html',total=5)
            self.assertIn('Download Excel',html)
            self.assertIn('Custom date range',html)
            empty=app.render_template('dashboard.html',total=0)
            self.assertIn('Import your first file',empty)
        with app.app.test_request_context('/upload'):
            self.assertIn('Import policy data',app.render_template('upload.html',total=0,last_upload=None,last_employee=None,result=None,employee_result=None))

if __name__=='__main__': unittest.main()
